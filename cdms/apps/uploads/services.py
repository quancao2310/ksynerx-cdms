import hashlib
import io
import os
import uuid
from datetime import timedelta
from itertools import islice
from pathlib import Path

from apps.products.constants import ChangeStatus
from apps.products.models import ProductVersion
from apps.products.services import apply_change
from django.conf import settings
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook

from .constants import (
    BOOL_COLUMNS,
    KNOWN_COLUMNS,
    LEASE_SECONDS,
    MAX_STORED_ERRORS,
    REQUIRED_COLUMNS,
    TEXT_COLUMNS,
)
from .models import ExcelUpload


class InvalidWorkbook(Exception):
    """The file is not a readable .xlsx or its header is missing required columns"""


class UploadBusy(Exception):
    """The file is currently being processed by another request"""


class ExcelUploadService:
    @staticmethod
    def is_blank(value):
        return value is None or (isinstance(value, str) and not value.strip())

    @staticmethod
    def number_to_text(value):
        if isinstance(value, bool):
            return value
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, (int, float)):
            return str(value)
        return value

    @staticmethod
    def to_int(value):
        if isinstance(value, bool):
            return value
        if isinstance(value, float) and value.is_integer():
            return int(value)
        if isinstance(value, str):
            text = value.strip()
            if text.isdigit():
                return int(text)
        return value

    @staticmethod
    def to_bool(value):
        if isinstance(value, bool):
            return value
        if isinstance(value, str):
            text = value.strip().lower()
            if text == "true":
                return True
            if text == "false":
                return False
        return value

    @classmethod
    def split(cls, value):
        if cls.is_blank(value):
            return []
        return [part.strip() for part in str(value).split(";") if part.strip()]

    @classmethod
    def row_to_raw(cls, header, values) -> dict:
        cells = dict(zip(header, values))
        raw = {name: None if cls.is_blank(cells.get(name)) else cells[name] for name in KNOWN_COLUMNS}

        raw["productId"] = cls.to_int(raw["productId"])
        for name in TEXT_COLUMNS:
            raw[name] = cls.number_to_text(raw[name])
        for name in BOOL_COLUMNS:
            raw[name] = cls.to_bool(raw[name])

        # units="box;pcs"
        raw["units"] = cls.split(raw["units"])

        # categories="CODE:Tên;CODE:Tên"
        categories = []
        for part in cls.split(raw["categories"]):
            code, _, name = part.partition(":")
            categories.append({
                "categoryCode": code.strip() or None,
                "categoryName": name.strip() or None,
            })
        raw["categories"] = categories
        return raw

    @staticmethod
    def read_header(rows):
        try:
            first = next(rows)
        except StopIteration:
            raise InvalidWorkbook("The file has no header row")
        header = [str(v).strip() if v is not None else "" for v in first]
        missing = [c for c in REQUIRED_COLUMNS if c not in header]
        if missing:
            raise InvalidWorkbook(f"Missing required columns: {', '.join(missing)}")
        return header

    @staticmethod
    def open_rows(source):
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
        except Exception:
            raise InvalidWorkbook("The file is not a valid .xlsx workbook")
        return workbook, workbook.active.iter_rows(values_only=True)

    @classmethod
    def validate_workbook(cls, content):
        workbook, rows = cls.open_rows(io.BytesIO(content))
        try:
            cls.read_header(rows)
            return sum(1 for _ in rows)
        finally:
            workbook.close()

    @staticmethod
    def store_file(sha256, content):
        directory = Path(settings.MEDIA_ROOT) / "uploads"
        directory.mkdir(parents=True, exist_ok=True)
        final = directory / f"{sha256}.xlsx"
        tmp = directory / f".{uuid.uuid4().hex}.tmp"
        with open(tmp, "wb") as fh:
            fh.write(content)
            fh.flush()
            os.fsync(fh.fileno())
        os.replace(tmp, final)
        return str(final)

    @classmethod
    def upload(cls, filename: str, content: bytes):
        total_rows = cls.validate_workbook(content)
        sha256 = hashlib.sha256(content).hexdigest()
        path = cls.store_file(sha256, content)

        try:
            with transaction.atomic():
                upload = ExcelUpload.objects.create(
                    sha256=sha256,
                    filename=filename,
                    file_path=path,
                    source_ts=timezone.now(),
                    total_rows=total_rows,
                )
            return upload, True
        except IntegrityError:
            return ExcelUpload.objects.get(sha256=sha256), False

    @staticmethod
    def lease_deadline():
        return timezone.now() + timedelta(seconds=LEASE_SECONDS)

    @classmethod
    def process(cls, upload_id):
        now = timezone.now()
        with transaction.atomic():
            upload = ExcelUpload.objects.select_for_update().get(pk=upload_id)
            if upload.status == ExcelUpload.Status.SUCCEEDED:
                return upload
            if upload.status == ExcelUpload.Status.PROCESSING and upload.lease_until and upload.lease_until > now:
                raise UploadBusy()
            upload.status = ExcelUpload.Status.PROCESSING
            upload.lease_until = cls.lease_deadline()
            upload.error = ""
            upload.save(update_fields=["status", "lease_until", "error"])

        chunk_size = settings.CDMS_UPLOAD_CHUNK_SIZE
        workbook = None
        try:
            workbook, rows = cls.open_rows(upload.file_path)
            header = cls.read_header(rows)
            rows = islice(rows, upload.processed_rows, None)

            counters = {s.value: getattr(upload, s.value) for s in ChangeStatus}
            errors = list(upload.errors)
            done = upload.processed_rows

            while True:
                chunk = list(islice(rows, chunk_size))
                if not chunk:
                    break
                for values in chunk:
                    done += 1
                    if all(cls.is_blank(v) for v in values):
                        continue
                    result = apply_change(
                        ProductVersion.Source.EXCEL,
                        cls.row_to_raw(header, values),
                        upload.source_ts,
                        source_ref=f"upload:{upload.pk}:row:{done + 1}",
                    )
                    counters[result.status.value] += 1
                    if result.errors and len(errors) < MAX_STORED_ERRORS:
                        errors.append({"row": done + 1, "errors": list(result.errors)})

                ExcelUpload.objects.filter(pk=upload.pk).update(
                    processed_rows=done, errors=errors, **counters,
                    lease_until=cls.lease_deadline(),
                )

            ExcelUpload.objects.filter(pk=upload.pk).update(
                status=ExcelUpload.Status.SUCCEEDED,
                finished_at=timezone.now(),
                lease_until=None,
                processed_rows=done,
                total_rows=done,
            )
        except Exception as exc:
            try:
                ExcelUpload.objects.filter(pk=upload.pk).update(
                    status=ExcelUpload.Status.FAILED,
                    lease_until=None,
                    error=str(exc),
                )
            except Exception:
                pass
            raise
        finally:
            if workbook is not None:
                workbook.close()

        return ExcelUpload.objects.get(pk=upload.pk)
